import csv
import io
import os
from datetime import datetime

from django.http import HttpResponse
from django.contrib.auth.models import User
from django.db.models import Count, Q
from django.utils.dateparse import parse_date
from django.conf import settings
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAdminUser
from rest_framework.response import Response

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor
from openpyxl.utils.units import pixels_to_EMU

from base.models import Room, Topic, Message


def _get_logo_path():
    """Finds the most recently updated logo image path across backend static and frontend assets."""
    candidates = [
        os.path.join(settings.BASE_DIR, '..', 'frontend', 'public', 'images', 'logo.png'),
        os.path.join(settings.BASE_DIR, 'base', 'static', 'images', 'logo.png'),
    ]
    existing = [p for p in candidates if os.path.exists(p)]
    if not existing:
        return None

    # Pick the most recently modified logo
    latest = max(existing, key=os.path.getmtime)

    # Sync to older copy so both frontend and backend stay consistent
    try:
        import shutil
        for other in existing:
            if other != latest and os.path.getmtime(other) < os.path.getmtime(latest):
                shutil.copy2(latest, other)
    except Exception:
        pass

    return latest


def _build_csv_response(filename, title, metadata_items, headers, rows):
    """Generates a CSV HttpResponse with metadata headers and UTF-8 BOM for Excel compatibility."""
    buffer = io.StringIO()
    buffer.write('\ufeff')
    writer = csv.writer(buffer)

    # Metadata Comment Headers
    writer.writerow([f"# STUDYBUD - {title.upper()}"])
    for key, val in metadata_items:
        writer.writerow([f"# {key}: {val}"])
    writer.writerow(["#"])

    # Table Column Headers and Data
    writer.writerow(headers)
    for row in rows:
        writer.writerow(row)

    response = HttpResponse(buffer.getvalue(), content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="{filename}.csv"'
    return response


def _build_excel_response(filename, title, metadata_items, headers, rows, sheet_title='Report', file_ext='xlsx'):
    """
    Generates a branded Excel spreadsheet:
    - Top centered StudyBud logo
    - Report title banner
    - Formatted metadata details block
    - Styled data table with navy headers, borders, and auto column widths
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    num_cols = max(len(headers), 6)

    # 1. Determine layout row positioning based on logo presence
    logo_path = _get_logo_path()
    if logo_path:
        ws.row_dimensions[1].height = 48
        ws.row_dimensions[2].height = 8
        ws.row_dimensions[3].height = 8
        title_row = 4
    else:
        title_row = 1

    # 2. Report Title Banner (Centered across all columns)
    ws.cell(row=title_row, column=1, value=f"STUDYBUD — {title.upper()}")
    ws.merge_cells(start_row=title_row, start_column=1, end_row=title_row, end_column=num_cols)
    title_cell = ws.cell(row=title_row, column=1)
    title_cell.font = Font(name='Calibri', size=14, bold=True, color='1E293B')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[title_row].height = 28

    # 3. Integrated Full-Width Executive Card Container
    card_border_color = 'CBD5E1'
    card_header_fill = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
    card_body_fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')

    meta_card_start_row = title_row + 2
    ws.row_dimensions[title_row + 1].height = 10
    ws.row_dimensions[meta_card_start_row].height = 24

    # Top Header Bar of the Card
    ws.cell(row=meta_card_start_row, column=1, value="REPORT DETAILS & FILTER CRITERIA")
    ws.merge_cells(start_row=meta_card_start_row, start_column=1, end_row=meta_card_start_row, end_column=num_cols)
    meta_title_cell = ws.cell(row=meta_card_start_row, column=1)
    meta_title_cell.font = Font(name='Calibri', size=10, bold=True, color='475569')
    meta_title_cell.alignment = Alignment(horizontal='center', vertical='center')

    for c in range(1, num_cols + 1):
        cell = ws.cell(row=meta_card_start_row, column=c)
        cell.fill = card_header_fill
        cell.border = Border(
            top=Side(style='thin', color=card_border_color),
            bottom=Side(style='thin', color=card_border_color),
            left=Side(style='thin', color=card_border_color) if c == 1 else None,
            right=Side(style='thin', color=card_border_color) if c == num_cols else None,
        )

    # Column configuration for internal items (Option 1: centered block with balanced margins)
    if num_cols >= 9:
        lbl1_col = 3
        val1_start = 4
        val1_end = 5
        lbl2_col = 6
        val2_start = 7
        val2_end = 8
    elif num_cols >= 7:
        lbl1_col = 2
        val1_start = 3
        val1_end = 4
        lbl2_col = 5
        val2_start = 6
        val2_end = num_cols - 1
    else:
        lbl1_col = 1
        val1_start = 2
        val1_end = 2
        lbl2_col = 3
        val2_start = 4
        val2_end = num_cols

    # Separate out Total Records for the dedicated summary footer bar
    clean_meta_items = [item for item in metadata_items if item[0] != "Total Records"]
    num_meta_rows = (len(clean_meta_items) + 1) // 2
    start_meta_row = meta_card_start_row + 1

    for r_idx in range(num_meta_rows):
        current_meta_row = start_meta_row + r_idx
        ws.row_dimensions[current_meta_row].height = 22

        item1 = clean_meta_items[r_idx * 2]
        item2 = clean_meta_items[r_idx * 2 + 1] if (r_idx * 2 + 1) < len(clean_meta_items) else None

        # Fill background and outer card borders
        for c in range(1, num_cols + 1):
            cell = ws.cell(row=current_meta_row, column=c)
            cell.fill = card_body_fill
            cell.border = Border(
                left=Side(style='thin', color=card_border_color) if c == 1 else None,
                right=Side(style='thin', color=card_border_color) if c == num_cols else None,
            )

        # Item 1 (Left side)
        c1 = ws.cell(row=current_meta_row, column=lbl1_col, value=f"{item1[0]}:")
        c1.font = Font(name='Calibri', size=10, bold=True, color='475569')
        c1.alignment = Alignment(horizontal='left', vertical='center')

        if val1_start < val1_end:
            ws.merge_cells(start_row=current_meta_row, start_column=val1_start, end_row=current_meta_row, end_column=val1_end)
        c2 = ws.cell(row=current_meta_row, column=val1_start, value=str(item1[1]))
        c2.font = Font(name='Calibri', size=10, color='1E293B')
        c2.alignment = Alignment(horizontal='left', vertical='center')

        # Item 2 (Right side)
        if item2:
            c3 = ws.cell(row=current_meta_row, column=lbl2_col, value=f"{item2[0]}:")
            c3.font = Font(name='Calibri', size=10, bold=True, color='475569')
            c3.alignment = Alignment(horizontal='left', vertical='center')

            if val2_start < val2_end:
                ws.merge_cells(start_row=current_meta_row, start_column=val2_start, end_row=current_meta_row, end_column=val2_end)
            c4 = ws.cell(row=current_meta_row, column=val2_start, value=str(item2[1]))
            c4.font = Font(name='Calibri', size=10, color='1E293B')
            c4.alignment = Alignment(horizontal='left', vertical='center')

    # Card Bottom Summary Bar (Total Records)
    summary_row = start_meta_row + num_meta_rows
    ws.row_dimensions[summary_row].height = 24
    ws.cell(row=summary_row, column=1, value=f"TOTAL RECORDS MATCHING CRITERIA: {len(rows)}")
    ws.merge_cells(start_row=summary_row, start_column=1, end_row=summary_row, end_column=num_cols)
    summary_cell = ws.cell(row=summary_row, column=1)
    summary_cell.font = Font(name='Calibri', size=10, bold=True, color='1E293B')
    summary_cell.alignment = Alignment(horizontal='center', vertical='center')

    for c in range(1, num_cols + 1):
        cell = ws.cell(row=summary_row, column=c)
        cell.fill = card_header_fill
        cell.border = Border(
            top=Side(style='thin', color=card_border_color),
            bottom=Side(style='thin', color=card_border_color),
            left=Side(style='thin', color=card_border_color) if c == 1 else None,
            right=Side(style='thin', color=card_border_color) if c == num_cols else None,
        )

    # 4. Spacing row before Table
    table_header_row = summary_row + 2
    ws.row_dimensions[table_header_row - 1].height = 14

    # 5. Data Table Header Row
    header_fill = PatternFill(start_color='2C3E50', end_color='2C3E50', fill_type='solid')
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    table_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1'),
    )

    ws.row_dimensions[table_header_row].height = 28
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=table_header_row, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = table_border

    # 6. Data Rows
    data_font = Font(name='Calibri', size=10)
    center_headers = {'id', 'room id', 'status', 'date joined', 'last login', 'created at', 'updated at',
                      'rooms hosted', 'rooms joined', 'messages sent', 'participants count', 'messages count'}

    for r_offset, row_data in enumerate(rows, start=1):
        r_num = table_header_row + r_offset
        ws.row_dimensions[r_num].height = 20
        is_even = r_offset % 2 == 0
        row_fill = PatternFill(start_color='F8F9FA', end_color='F8F9FA', fill_type='solid') if is_even else None

        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=r_num, column=col_idx, value=val)
            cell.font = data_font
            cell.border = table_border

            header_name = str(headers[col_idx - 1]).strip().lower()
            if header_name in center_headers:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')

            if row_fill:
                cell.fill = row_fill

    # 7. Auto-adjust column widths based on headers and data
    col_pixel_widths = []
    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = len(str(headers[col_idx - 1]))
        for r in range(table_header_row + 1, table_header_row + len(rows) + 1):
            cell_val = ws.cell(row=r, column=col_idx).value
            if cell_val is not None:
                max_len = max(max_len, len(str(cell_val)))
        col_w = max(max_len + 4, 12)
        ws.column_dimensions[col_letter].width = col_w
        # Excel column width conversion: 1 width unit ≈ 7.5 px + 5 px padding
        col_pixel_widths.append((col_idx, int(col_w * 7.5 + 5)))

    # 8. Accurately center the Logo at Row 1 directly above the title banner
    if logo_path:
        try:
            img = OpenpyxlImage(logo_path)
            target_height = 42
            if img.height and img.height > 0:
                aspect_ratio = img.width / img.height
                img.height = target_height
                img.width = max(24, int(target_height * aspect_ratio))
            else:
                img.width = 130
                img.height = 42

            total_px = sum(px for _, px in col_pixel_widths[:num_cols])
            center_x = total_px / 2
            logo_start_x = max(0, center_x - (img.width / 2))

            curr_x = 0
            anchor_col = 1
            anchor_offset = 0
            for c_idx, px in col_pixel_widths[:num_cols]:
                if curr_x + px > logo_start_x:
                    anchor_col = c_idx
                    anchor_offset = int(logo_start_x - curr_x)
                    break
                curr_x += px

            anchor = OneCellAnchor()
            anchor._from.col = anchor_col - 1
            anchor._from.colOff = pixels_to_EMU(anchor_offset)
            anchor._from.row = 0
            anchor._from.rowOff = pixels_to_EMU(3)
            anchor.ext.width = pixels_to_EMU(img.width)
            anchor.ext.height = pixels_to_EMU(img.height)
            img.anchor = anchor
            ws.add_image(img)
        except Exception:
            pass

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    content_type = (
        'application/vnd.ms-excel' if file_ext == 'xls'
        else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response = HttpResponse(buf.getvalue(), content_type=content_type)
    response['Content-Disposition'] = f'attachment; filename="{filename}.{file_ext}"'
    return response


@api_view(['GET'])
@permission_classes([IsAdminUser])
def admin_report_filter_options(request):
    """Provides dropdown filter options (users, roles, topics) for reports."""
    users = (
        User.objects.all()
        .order_by('username')
        .values('id', 'username', 'first_name', 'last_name', 'email')
    )
    topics = (
        Topic.objects.all()
        .order_by('name')
        .values('id', 'name')
    )

    roles = [
        {'value': 'all', 'label': 'All Roles'},
        {'value': 'superuser', 'label': 'Admin / Superuser'},
        {'value': 'staff', 'label': 'Staff'},
        {'value': 'member', 'label': 'Regular Member'},
        {'value': 'active', 'label': 'Active Accounts'},
        {'value': 'inactive', 'label': 'Inactive / Banned'},
    ]

    return Response({
        'users': list(users),
        'topics': list(topics),
        'roles': roles,
    })


@api_view(['GET'])
@permission_classes([IsAdminUser])
def admin_user_report(request):
    """
    Generates User Report filtered by role, user_id, or search.
    Exports include app logo branding and metadata summary.
    If no matching records, returns 404 No data found.
    """
    role = request.GET.get('role', 'all')
    user_id = request.GET.get('user_id')
    search = request.GET.get('search', '').strip()
    export_format = (request.GET.get('export') or request.GET.get('file_format') or '').lower()

    qs = (
        User.objects.all()
        .select_related('profile')
        .annotate(
            rooms_hosted_count=Count('room', distinct=True),
            messages_sent_count=Count('message', distinct=True),
            rooms_joined_count=Count('participants', distinct=True),
        )
        .order_by('-date_joined')
    )

    role_label = 'All Roles'
    if role == 'superuser':
        qs = qs.filter(is_superuser=True)
        role_label = 'Admin / Superuser'
    elif role == 'staff':
        qs = qs.filter(is_staff=True, is_superuser=False)
        role_label = 'Staff'
    elif role == 'member':
        qs = qs.filter(is_staff=False, is_superuser=False)
        role_label = 'Regular Member'
    elif role == 'active':
        qs = qs.filter(is_active=True)
        role_label = 'Active Accounts'
    elif role == 'inactive':
        qs = qs.filter(is_active=False)
        role_label = 'Inactive Accounts'

    selected_user_label = 'All Users'
    if user_id and user_id != 'all':
        qs = qs.filter(id=user_id)
        u_obj = User.objects.filter(id=user_id).first()
        if u_obj:
            selected_user_label = u_obj.username

    if search:
        qs = qs.filter(Q(username__icontains=search) | Q(email__icontains=search))

    user_list = list(qs)

    # Handle Export Formats
    if export_format in ['csv', 'xls', 'xlsx']:
        if not user_list:
            return Response({'detail': 'No data found matching the selected filters.'}, status=404)

        timestamp_str = datetime.now().strftime('%Y%m%d_%H%M%S')
        readable_time = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        filename = f"user_report_{timestamp_str}"
        title = "User Activity Report"

        metadata_items = [
            ("Report Name", title),
            ("Generated At", readable_time),
            ("Role Filter", role_label),
            ("User Filter", selected_user_label),
            ("Total Records", len(user_list)),
        ]
        if search:
            metadata_items.append(("Search Query", search))

        headers = [
            'ID',
            'Username',
            'Email',
            'Role',
            'Status',
            'Date Joined',
            'Last Login',
            'Rooms Hosted',
            'Rooms Joined',
            'Messages Sent',
        ]
        rows = []
        for u in user_list:
            r_label = 'Superuser' if u.is_superuser else ('Staff' if u.is_staff else 'Member')
            status_label = 'Active' if u.is_active else 'Inactive'
            joined_str = u.date_joined.strftime('%Y-%m-%d %H:%M') if u.date_joined else ''
            login_str = u.last_login.strftime('%Y-%m-%d %H:%M') if u.last_login else 'Never'
            rows.append([
                u.id,
                u.username,
                u.email or 'N/A',
                r_label,
                status_label,
                joined_str,
                login_str,
                u.rooms_hosted_count,
                u.rooms_joined_count,
                u.messages_sent_count,
            ])

        if export_format == 'csv':
            return _build_csv_response(filename, title, metadata_items, headers, rows)
        else:
            return _build_excel_response(filename, title, metadata_items, headers, rows, sheet_title='User Report', file_ext=export_format)

    return Response({
        'summary': {
            'total_users': len(user_list),
            'total_rooms_hosted': sum(u.rooms_hosted_count for u in user_list),
            'total_messages_sent': sum(u.messages_sent_count for u in user_list),
        },
        'results': [],
    })


@api_view(['GET'])
@permission_classes([IsAdminUser])
def admin_room_report(request):
    """
    Generates Room Report.
    Date range (start_date and end_date) is strictly mandatory.
    Exports include app logo branding and metadata summary.
    If no matching records, returns 404 No data found.
    """
    topic_id = request.GET.get('topic_id')
    creator_id = request.GET.get('creator_id')
    participant_id = request.GET.get('participant_id')
    start_date = request.GET.get('start_date')
    end_date = request.GET.get('end_date')
    search = request.GET.get('search', '').strip()
    export_format = (request.GET.get('export') or request.GET.get('file_format') or '').lower()

    # Mandatory Date Range Validation
    if not start_date or not end_date:
        return Response({'detail': 'Both start_date and end_date are required.'}, status=400)

    parsed_start = parse_date(start_date)
    parsed_end = parse_date(end_date)
    if not parsed_start or not parsed_end:
        return Response({'detail': 'Invalid date format. Use YYYY-MM-DD.'}, status=400)

    if parsed_start > parsed_end:
        return Response({'detail': 'From Date cannot be later than To Date.'}, status=400)

    qs = (
        Room.objects.all()
        .select_related('host', 'topic')
        .prefetch_related('participants')
        .annotate(
            participant_count=Count('participants', distinct=True),
            message_count=Count('message', distinct=True),
        )
        .filter(
            created__date__gte=parsed_start,
            created__date__lte=parsed_end
        )
        .order_by('-created')
    )

    topic_label = 'All Topics'
    if topic_id and topic_id != 'all':
        qs = qs.filter(topic_id=topic_id)
        t_obj = Topic.objects.filter(id=topic_id).first()
        if t_obj:
            topic_label = t_obj.name

    creator_label = 'All Creators'
    if creator_id and creator_id != 'all':
        qs = qs.filter(host_id=creator_id)
        c_obj = User.objects.filter(id=creator_id).first()
        if c_obj:
            creator_label = c_obj.username

    participant_label = 'All Participants'
    if participant_id and participant_id != 'all':
        qs = qs.filter(participants__id=participant_id)
        p_obj = User.objects.filter(id=participant_id).first()
        if p_obj:
            participant_label = p_obj.username

    if search:
        qs = qs.filter(Q(name__icontains=search) | Q(description__icontains=search))

    room_list = list(qs)

    # Handle Export Formats
    if export_format in ['csv', 'xls', 'xlsx']:
        if not room_list:
            return Response({'detail': 'No data found matching the selected filters.'}, status=404)

        timestamp_str = datetime.now().strftime('%Y%m%d_%H%M%S')
        readable_time = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        filename = f"room_report_{timestamp_str}"
        title = "Room Activity Report"

        metadata_items = [
            ("Report Name", title),
            ("Generated At", readable_time),
            ("Date Range", f"{start_date} to {end_date}"),
            ("Topic Filter", topic_label),
            ("Creator (Host)", creator_label),
            ("Participant", participant_label),
            ("Total Records", len(room_list)),
        ]
        if search:
            metadata_items.append(("Search Query", search))

        headers = [
            'Room ID',
            'Room Name',
            'Topic',
            'Creator (Host)',
            'Date Created',
            'Last Updated',
            'Participants Count',
            'Messages Count',
            'Participants List',
        ]
        rows = []
        for r in room_list:
            host_str = r.host.username if r.host else 'Deleted User'
            t_name = r.topic.name if r.topic else 'General'
            created_str = r.created.strftime('%Y-%m-%d %H:%M') if r.created else ''
            updated_str = r.updated.strftime('%Y-%m-%d %H:%M') if r.updated else ''
            participants_str = ', '.join([p.username for p in r.participants.all()[:15]])
            rows.append([
                r.id,
                r.name,
                t_name,
                host_str,
                created_str,
                updated_str,
                r.participant_count,
                r.message_count,
                participants_str,
            ])

        if export_format == 'csv':
            return _build_csv_response(filename, title, metadata_items, headers, rows)
        else:
            return _build_excel_response(filename, title, metadata_items, headers, rows, sheet_title='Room Report', file_ext=export_format)

    return Response({
        'summary': {
            'total_rooms': len(room_list),
            'total_messages': sum(r.message_count for r in room_list),
            'total_participants': sum(r.participant_count for r in room_list),
        },
        'results': [],
    })
