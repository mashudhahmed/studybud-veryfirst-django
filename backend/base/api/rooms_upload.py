import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from django.http import HttpResponse
from django.contrib.auth.models import User
from django.db import transaction
from rest_framework.decorators import api_view, permission_classes
from rest_framework.response import Response
from rest_framework import status

from base.api.permissions import IsSuperUser
from base.models import Room, Topic


@api_view(['GET'])
@permission_classes([IsSuperUser])
def admin_rooms_upload_template(request):
    """
    Generates and downloads an Excel (.xlsx) template for bulk uploading rooms.
    Includes styled headers and sample rows to guide administrators.
    Room host is automatically assigned to the uploading administrator.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rooms Import Template"

    headers = ["name", "topic", "description", "participants"]
    sample_rows = [
        [
            "Fullstack React & Django Mastery",
            "Python",
            "Discussions on React 19, Django REST Framework, and architecture.",
            "alice, bob",
        ],
        [
            "Algorithms & LeetCode Study Group",
            "Algorithms",
            "Weekly competitive programming and technical interview prep.",
            "",
        ],
        [
            "Cloud & DevOps Fundamentals",
            "DevOps",
            "Docker, Kubernetes, and CI/CD pipelines.",
            "alice",
        ],
    ]

    # Header styling
    header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center")
    border_style = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    ws.row_dimensions[1].height = 26
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = border_style

    # Data rows
    data_font = Font(name="Calibri", size=10)
    for row_idx, row_data in enumerate(sample_rows, start=2):
        ws.row_dimensions[row_idx].height = 20
        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = border_style
            cell.alignment = Alignment(horizontal="left", vertical="center")

    # Column width auto-fit
    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = len(headers[col_idx - 1])
        for row in sample_rows:
            val_len = len(str(row[col_idx - 1]))
            if val_len > max_len:
                max_len = val_len
        ws.column_dimensions[col_letter].width = max(max_len + 5, 14)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    response = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="studybud_rooms_template.xlsx"'
    return response


@api_view(['POST'])
@permission_classes([IsSuperUser])
def admin_bulk_upload_rooms(request):
    """
    Parses an uploaded Excel (.xlsx) file and creates rooms in bulk using an
    All-or-Nothing (Atomic) policy:
      - Phase 1: In-memory dry-run validation of all rows (duplicates, format).
      - Host is automatically assigned to the uploading admin (request.user).
      - If ANY error or duplicate exists, the entire batch is aborted (0 rooms/topics created).
      - Phase 2: If 100% valid, all rooms and topics are committed in a single atomic transaction.
    """
    uploaded_file = request.FILES.get('file')
    if not uploaded_file:
        return Response(
            {'detail': 'No file uploaded. Please choose an Excel file.'},
            status=status.HTTP_400_BAD_REQUEST,
        )

    if not uploaded_file.name.lower().endswith('.xlsx'):
        return Response(
            {'detail': 'Invalid file format. Please upload an Excel (.xlsx) file.'},
            status=status.HTTP_400_BAD_REQUEST,
        )

    if uploaded_file.size > 10 * 1024 * 1024:
        return Response(
            {'detail': 'File size exceeds 10MB limit.'},
            status=status.HTTP_400_BAD_REQUEST,
        )

    try:
        wb = openpyxl.load_workbook(uploaded_file, data_only=True)
        ws = wb.active
    except Exception as err:
        return Response(
            {'detail': f'Unable to read Excel file: {str(err)}'},
            status=status.HTTP_400_BAD_REQUEST,
        )

    all_rows = list(ws.iter_rows(values_only=True))
    if not all_rows:
        return Response(
            {'detail': 'The uploaded Excel file is empty.'},
            status=status.HTTP_400_BAD_REQUEST,
        )

    # Aliases for flexible column matching
    NAME_ALIASES = {'name', 'room', 'room_name', 'room name', 'title', 'room title'}
    TOPIC_ALIASES = {'topic', 'topic_name', 'topic name', 'category', 'subject'}
    DESC_ALIASES = {'description', 'desc', 'details', 'about', 'body'}
    PARTS_ALIASES = {'participants', 'members', 'users', 'participant', 'member'}

    # Scan the first 10 rows to locate header row
    header_row_idx = None
    header_map = {}
    for r_idx, row in enumerate(all_rows[:10]):
        if not row:
            continue
        temp_map = {}
        for c_idx, cell in enumerate(row):
            if cell is not None:
                cleaned = str(cell).strip().lower().replace('_', ' ')
                raw_clean = str(cell).strip().lower()
                if raw_clean in NAME_ALIASES or cleaned in NAME_ALIASES:
                    temp_map['name'] = c_idx
                elif raw_clean in TOPIC_ALIASES or cleaned in TOPIC_ALIASES:
                    temp_map['topic'] = c_idx
                elif raw_clean in DESC_ALIASES or cleaned in DESC_ALIASES:
                    temp_map['description'] = c_idx
                elif raw_clean in PARTS_ALIASES or cleaned in PARTS_ALIASES:
                    temp_map['participants'] = c_idx
        if 'name' in temp_map:
            header_row_idx = r_idx
            header_map = temp_map
            break

    if header_row_idx is None or 'name' not in header_map:
        return Response(
            {'detail': 'Excel sheet must contain a "name" or "room" column header.'},
            status=status.HTTP_400_BAD_REQUEST,
        )

    name_idx = header_map['name']
    topic_idx = header_map.get('topic')
    desc_idx = header_map.get('description')
    parts_idx = header_map.get('participants')

    # Pre-fetch existing database state for O(1) in-memory lookups
    existing_rooms_lower = {
        name.lower(): name for name in Room.objects.values_list('name', flat=True)
    }
    seen_in_batch = set()

    topics_cache = {t.name.lower(): t for t in Topic.objects.all()}

    all_users = list(User.objects.all())
    user_by_username = {u.username.lower(): u for u in all_users}
    user_by_email = {u.email.lower(): u for u in all_users if u.email}
    user_by_id = {u.id: u for u in all_users}

    def resolve_user(identifier):
        if identifier is None:
            return None
        s_val = str(identifier).strip()
        if not s_val:
            return None
        try:
            val_float = float(s_val)
            if val_float.is_integer():
                uid = int(val_float)
                if uid in user_by_id:
                    return user_by_id[uid]
        except ValueError:
            pass

        lower_val = s_val.lower()
        if lower_val.startswith('@'):
            lower_val = lower_val[1:]

        if lower_val in user_by_username:
            return user_by_username[lower_val]
        if lower_val in user_by_email:
            return user_by_email[lower_val]
        return None

    data_rows = all_rows[header_row_idx + 1:]
    parsed_rows = []
    duplicates = []
    errors = []
    total_processed = 0

    # The room host is always the administrator who uploaded the spreadsheet
    uploading_admin = request.user

    # Phase 1: Validate all rows (Dry-run)
    for offset, row in enumerate(data_rows):
        row_num = header_row_idx + 2 + offset
        # Skip completely empty rows
        if not row or all(c is None or str(c).strip() == '' for c in row):
            continue

        total_processed += 1

        raw_name = str(row[name_idx]).strip() if name_idx < len(row) and row[name_idx] is not None else ''

        # Validation: Room name required & minimum 3 characters
        if not raw_name or len(raw_name) < 3:
            errors.append({
                'row': row_num,
                'room_name': raw_name or '(Empty)',
                'error': 'Room name is required and must be at least 3 characters long.',
            })
            continue

        name_lower = raw_name.lower()

        # DUPLICATE CHECK: In database
        if name_lower in existing_rooms_lower:
            duplicates.append({
                'row': row_num,
                'room_name': raw_name,
                'reason': f"A room named '{existing_rooms_lower[name_lower]}' already exists in StudyBud.",
            })
            continue

        # DUPLICATE CHECK: Intra-file duplicate
        if name_lower in seen_in_batch:
            duplicates.append({
                'row': row_num,
                'room_name': raw_name,
                'reason': 'Duplicate room name found earlier in this spreadsheet.',
            })
            continue

        seen_in_batch.add(name_lower)

        # Topic string
        raw_topic = ''
        if topic_idx is not None and topic_idx < len(row) and row[topic_idx] is not None:
            raw_topic = str(row[topic_idx]).strip()

        # Description
        description = ''
        if desc_idx is not None and desc_idx < len(row) and row[desc_idx] is not None:
            description = str(row[desc_idx]).strip()

        # Participants resolution (uploading admin is automatically a participant)
        participants_set = {uploading_admin}
        if parts_idx is not None and parts_idx < len(row) and row[parts_idx] is not None:
            raw_parts = str(row[parts_idx]).strip()
            if raw_parts:
                tokens = [t.strip() for t in raw_parts.replace(';', ',').split(',') if t.strip()]
                for token in tokens:
                    part_user = resolve_user(token)
                    if part_user:
                        participants_set.add(part_user)

        parsed_rows.append({
            'name': raw_name,
            'raw_topic': raw_topic,
            'description': description,
            'host': uploading_admin,
            'participants': list(participants_set),
        })

    if total_processed == 0:
        return Response(
            {'detail': 'The spreadsheet does not contain any data rows.'},
            status=status.HTTP_400_BAD_REQUEST,
        )

    # ALL-OR-NOTHING GATE: If ANY duplicates or errors occurred, abort entire upload (0 created)
    if duplicates or errors:
        return Response(
            {
                'total_rows': total_processed,
                'created_count': 0,
                'skipped_duplicates_count': len(duplicates),
                'duplicates': duplicates,
                'errors': errors,
            },
            status=status.HTTP_200_OK,
        )

    # Phase 2: Atomic Execution (Only runs if 100% of rows passed)
    created_rooms = []
    try:
        with transaction.atomic():
            for item in parsed_rows:
                topic_obj = None
                raw_topic = item['raw_topic']
                if raw_topic:
                    t_lower = raw_topic.lower()
                    if t_lower in topics_cache:
                        topic_obj = topics_cache[t_lower]
                    else:
                        topic_obj = Topic.objects.create(name=raw_topic)
                        topics_cache[t_lower] = topic_obj

                room = Room.objects.create(
                    name=item['name'],
                    topic=topic_obj,
                    host=item['host'],
                    description=item['description'],
                )
                if item['participants']:
                    room.participants.set(item['participants'])
                created_rooms.append({'id': room.id, 'name': room.name})
    except Exception as exc:
        return Response(
            {'detail': f'Failed to create rooms: {str(exc)}'},
            status=status.HTTP_500_INTERNAL_SERVER_ERROR,
        )

    return Response(
        {
            'total_rows': total_processed,
            'created_count': len(created_rooms),
            'skipped_duplicates_count': 0,
            'duplicates': [],
            'errors': [],
        },
        status=status.HTTP_200_OK,
    )
